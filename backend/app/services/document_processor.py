"""Document processing service."""

from __future__ import annotations

import base64
from io import BytesIO
from typing import Any

import pandas as pd
import pytesseract
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader
from openai import AsyncOpenAI
from PIL import Image
from pypdf import PdfReader
from docx import Document as DocxDocument

from app.config import settings
from app.utils.logger import get_logger

logger = get_logger()


class DocumentProcessor:
	"""Handles parsing of PDF, DOCX, TXT, images, CSV, XLSX files."""

	def __init__(self) -> None:
		self._client = AsyncOpenAI(api_key=settings.OPENAI_API_KEY)

	async def process(self, file_path: str, file_type: str, doc_id: str) -> list[Document]:
		"""Parse a document and split it into chunks."""
		if file_type == "pdf":
			docs = self._load_pdf(file_path)
		elif file_type == "docx":
			docs = self._load_docx(file_path)
		elif file_type in {"png", "jpg", "jpeg"}:
			docs = await self._load_image(file_path)
		elif file_type == "csv":
			docs = self._load_csv(file_path)
		elif file_type == "xlsx":
			docs = self._load_xlsx(file_path)
		elif file_type == "txt":
			docs = self._load_txt(file_path)
		else:
			raise ValueError(f"Unsupported file type: {file_type}")

		for doc in docs:
			doc.metadata["doc_id"] = doc_id
			doc.metadata["file_type"] = file_type

		return self._split_documents(docs)

	def _load_pdf(self, path: str) -> list[Document]:
		loader = PyPDFLoader(path)
		docs = loader.load()
		images_by_page = self._extract_pdf_images(path)

		for doc in docs:
			page = int(doc.metadata.get("page", 0))
			if not doc.page_content.strip() and page in images_by_page:
				ocr_text = pytesseract.image_to_string(images_by_page[page])
				doc.page_content = ocr_text
			if page in images_by_page:
				encoded = self._encode_image(images_by_page[page])
				doc.metadata["image_base64"] = encoded
			doc.metadata["source"] = path
		return docs

	def _extract_pdf_images(self, path: str) -> dict[int, Image.Image]:
		images: dict[int, Image.Image] = {}
		reader = PdfReader(path)
		for idx, page in enumerate(reader.pages):
			try:
				page_images = list(getattr(page, "images", []))
			except Exception:
				page_images = []
			if page_images:
				image = page_images[0]
				data = image.data
				img = Image.open(BytesIO(data))
				images[idx] = img
		return images

	def _load_docx(self, path: str) -> list[Document]:
		doc = DocxDocument(path)
		documents: list[Document] = []
		for paragraph in doc.paragraphs:
			text = paragraph.text.strip()
			if text:
				level = paragraph.style.name if paragraph.style else "Normal"
				documents.append(Document(page_content=text, metadata={"source": path, "heading": level}))

		for table in doc.tables:
			rows = []
			for row in table.rows:
				rows.append([cell.text.strip() for cell in row.cells])
			if rows:
				headers = rows[0]
				body = rows[1:]
				markdown = self._to_markdown_table(headers, body)
				documents.append(Document(page_content=markdown, metadata={"source": path, "table": True}))

		return documents

	async def _load_image(self, path: str) -> list[Document]:
		image = Image.open(path)
		ocr_text = pytesseract.image_to_string(image)
		encoded = self._encode_image(image)

		description = await self._describe_image(encoded)
		content = f"OCR Text:\n{ocr_text}\n\nImage Description:\n{description}"
		return [Document(page_content=content, metadata={"source": path, "image_base64": encoded})]

	def _load_csv(self, path: str) -> list[Document]:
		df = pd.read_csv(path)
		documents: list[Document] = []
		for idx, row in df.iterrows():
			items = ", ".join([f"{col}={row[col]}" for col in df.columns])
			documents.append(Document(page_content=f"Row {idx + 1}: {items}", metadata={"source": path}))

		summary = df.describe(include="all").to_markdown()
		documents.append(Document(page_content=f"Summary:\n{summary}", metadata={"source": path, "summary": True}))
		return documents

	def _load_xlsx(self, path: str) -> list[Document]:
		data = pd.read_excel(path, sheet_name=None)
		documents: list[Document] = []
		for sheet_name, df in data.items():
			markdown = df.to_markdown(index=False)
			documents.append(
				Document(page_content=f"Sheet: {sheet_name}\n{markdown}", metadata={"source": path, "sheet": sheet_name})
			)
		return documents

	def _load_txt(self, path: str) -> list[Document]:
		with open(path, "r", encoding="utf-8") as handle:
			content = handle.read()
		return [Document(page_content=content, metadata={"source": path})]

	def _split_documents(self, docs: list[Document]) -> list[Document]:
		splitter = RecursiveCharacterTextSplitter(
			chunk_size=settings.CHUNK_SIZE,
			chunk_overlap=settings.CHUNK_OVERLAP,
		)
		chunks = splitter.split_documents(docs)
		for idx, chunk in enumerate(chunks):
			chunk.metadata["chunk_index"] = idx
		return chunks

	def _encode_image(self, image: Image.Image) -> str:
		buffer = BytesIO()
		image.save(buffer, format="PNG")
		return base64.b64encode(buffer.getvalue()).decode("utf-8")

	async def _describe_image(self, image_base64: str) -> str:
		response = await self._client.chat.completions.create(
			model=settings.OPENAI_CHAT_MODEL,
			messages=[
				{
					"role": "user",
					"content": [
						{"type": "text", "text": "Provide a detailed description of this image."},
						{
							"type": "image_url",
							"image_url": {"url": f"data:image/png;base64,{image_base64}"},
						},
					],
				}
			],
			max_tokens=512,
		)
		return response.choices[0].message.content or ""

	def _to_markdown_table(self, headers: list[str], rows: list[list[str]]) -> str:
		header_row = "| " + " | ".join(headers) + " |"
		divider = "| " + " | ".join(["---" for _ in headers]) + " |"
		body_rows = ["| " + " | ".join(row) + " |" for row in rows]
		return "\n".join([header_row, divider, *body_rows])
